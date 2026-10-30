import json, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from PIL import Image

FIG = "/home/user/BIO-AlGO/report/source/fig/"
F = json.load(open("/home/user/BIO-AlGO/report/source/facts.json"))
rgb = lambda h: RGBColor.from_string(h)
BLUE, ORANGE, GRAY, DARK, GREEN = "2683C6", "D9731A", "8A97A8", "335B74", "42BA97"
TINT, TINT2, LINE, WHITE, MID = "EAF3FA", "F3F6F8", "D5DEE5", "FFFFFF", "9CC8E8"

prs = Presentation("work.pptx")
S = list(prs.slides)


# ------------------------------------------------------------------ helpers
def drop_subtitle(slide):
    for sh in list(slide.placeholders):
        if sh.placeholder_format.type is not None and "SUBTITLE" in str(sh.placeholder_format.type):
            sh._element.getparent().remove(sh._element)

def set_title(slide, text, size=36):
    t = slide.shapes.title
    t.text_frame.text = text
    t.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    for r in t.text_frame.paragraphs[0].runs:
        r.font.size = Pt(size)

def para_fmt(p, size, bold=False, color=DARK, italic=False, align=PP_ALIGN.LEFT, font="Calibri"):
    for r in p.runs:
        r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
        r.font.color.rgb = rgb(color)
    p.alignment = align

def text(slide, x, y, w, h, s, size=16, bold=False, color=DARK, align="l", anchor="t", italic=False, name=None, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(line, tuple): txt, sz, b, c = line
        else: txt, sz, b, c = line, size, bold, color
        p.text = txt
        para_fmt(p, sz, b, c, italic, {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align])
        if spacing: p.space_after = Pt(spacing)
    return tb

def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name: sp.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sp.adjustments[0] = radius
    sp.shadow.inherit = False
    if fill: sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    else: sp.fill.background()
    if line: sp.line.color.rgb = rgb(line); sp.line.width = Pt(1)
    else: sp.line.fill.background()
    return sp

def label_box(slide, x, y, w, h, s, fill=TINT, color=DARK, size=14, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, align="c", line=None, anchor="m", name=None):
    sp = box(slide, x, y, w, h, fill, line, shape, radius, name)
    tf = sp.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
    lines = s if isinstance(s, list) else [s]
    for i, line_ in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(line_, tuple): txt, sz, b, c = line_
        else: txt, sz, b, c = line_, size, bold, color
        p.text = txt; para_fmt(p, sz, b, c, False, {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[align])
    return sp

def circle(slide, x, y, d, fill, s="", size=16, color=WHITE, bold=True):
    sp = label_box(slide, x, y, d, d, s, fill=fill, color=color, size=size, bold=bold, shape=MSO_SHAPE.OVAL)
    tf = sp.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.word_wrap = False
    return sp

def arrow(slide, x1, y1, x2, y2, color=GRAY):
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    from lxml import etree
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd"); tail.set("type", "triangle")
    return c

def pic(slide, path, x, y, w=None, h=None, name=None):
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)
    if name: p.name = name
    return p

def notes(slide, s): slide.notes_slide.notes_text_frame.text = s

def style_chart(ch, size=12, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM):
    ch.font.size = Pt(size); ch.font.name = "Calibri"; ch.font.color.rgb = rgb(DARK)
    ch.has_legend = legend
    if legend:
        ch.legend.position = legend_pos; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(size)
    va, ca = ch.value_axis, ch.category_axis
    va.major_gridlines.format.line.color.rgb = rgb(LINE); va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size); ca.tick_labels.font.size = Pt(size)
    ca.format.line.color.rgb = rgb(LINE)
    va.has_major_gridlines = True

def chart_title(ch, s, size=14):
    ch.has_title = True; ch.chart_title.text_frame.text = s
    r = ch.chart_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = rgb(DARK); r.font.name = "Calibri"
    ch.chart_title.include_in_layout = False

def color_series(ser, hexc):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(hexc)

def prep(slide, title):
    drop_subtitle(slide); set_title(slide, title)

# ------------------------------------------------------------------ 0 title
s = S[0]
drop_subtitle(s)
t = s.shapes.title
t.top, t.height, t.left, t.width = Inches(0.3), Inches(3.9), Inches(2.0), Inches(9.6)
tf = t.text_frame; tf.auto_size = MSO_AUTO_SIZE.NONE
runs = tf.paragraphs[0].runs
sizes = [40, 16, 24, 16, 22]
for r, sz in zip(runs, sizes):
    r.font.size = Pt(sz); r.font.name = 'Calibri'
runs[4].font.bold = True; runs[4].font.color.rgb = rgb(DARK)
runs[2].text = "Department of Artificial Intelligence and Data Science"
runs[4].text = "SwarmFolio: Multi-Objective Stock Portfolio Optimization Using Particle Swarm Optimization and Genetic Algorithms"
runs[3].text = "Project Work Phase-2 (23AI7PWPP2) Final Presentation on"
text(s, 2.3, 4.45, 4.7, 0.4, "Submitted by", 16, True, BLUE, name="Team heading")
people = [("Vivek Boora", "1BM23AD073"), ("Ayush Pandey", "1BM23AD074"), ("Nitin", "1BM24AD404")]
for i, (n, u) in enumerate(people):
    text(s, 2.3, 4.9 + i * 0.45, 2.6, 0.4, n, 18, True, DARK, name=f"Team name {i+1}")
    text(s, 4.9, 4.9 + i * 0.45, 2.2, 0.4, u, 18, False, DARK, name=f"Team USN {i+1}")
text(s, 7.9, 4.45, 4.6, 0.4, "Under the guidance of", 16, True, BLUE, name="Guide heading")
text(s, 7.9, 4.9, 4.6, 1.4, [("Prof. Suman M", 20, True, DARK), ("Assistant Professor", 18, False, DARK), ("Dept of AI & DS", 18, False, DARK)], name="Guide")
notes(s, "Introduce the project and the team. SwarmFolio compares a particle swarm optimiser and a genetic algorithm on a three-objective stock portfolio problem.")

# ------------------------------------------------------------------ 1 TOC
s = S[1]; prep(s, "Table of Contents")
items = [("Abstract", 3), ("Introduction", 4), ("Literature Survey", 6), ("Requirement Analysis and Specification", 7), ("Design", 8),
         ("Implementation", 10), ("Testing and Validation", 12), ("Results and Discussion", 13), ("Conclusion", 17), ("POs mapping matrix", 18)]
for i, (n, sl) in enumerate(items):
    col, row = divmod(i, 5)
    x, y = 0.95 + col * 6.1, 1.85 + row * 0.96
    box(s, x, y, 5.75, 0.8, TINT, name=f"TOC tile {i+1}")
    circle(s, x + 0.15, y + 0.1, 0.6, BLUE, str(sl), 18)
    text(s, x + 0.9, y + 0.1, 4.7, 0.6, n, 20, False, DARK, anchor="m", name=f"TOC item {i+1}")
text(s, 0.95, 6.65, 8, 0.3, "Circle shows the slide number", 12, False, GRAY)
notes(s, "Roadmap of the talk. The circles show the slide number where each section starts.")

# ------------------------------------------------------------------ 2 Abstract
s = S[2]; prep(s, "Abstract")
cols = [("Three goals", BLUE), ("Two optimisers", ORANGE), ("What we found", GREEN)]
xs = [0.65, 4.8, 8.95]
for (h, c), x in zip(cols, xs):
    box(s, x, 1.55, 3.85, 4.35, TINT2, name=f"Card {h}")
    label_box(s, x + 0.25, 1.75, 3.35, 0.6, h, fill=c, color=WHITE, size=20, bold=True, radius=0.3)
x = xs[0] + 0.3
for i, (lab, c) in enumerate([("Maximise return", BLUE), ("Minimise volatility", DARK), ("Minimise CVaR (tail loss)", GRAY)]):
    label_box(s, x, 2.65 + i * 0.85, 3.25, 0.65, lab, fill=WHITE, color=DARK, size=18, line=c, radius=0.3)
text(s, x, 5.2, 3.3, 0.5, "under cap, floor and holding limits", 14, False, GRAY, "c")
x = xs[1] + 0.3
label_box(s, x, 2.65, 3.25, 0.8, ["Particle swarm", ("MOPSO", 14, False, GRAY)], fill=WHITE, line=BLUE, size=18, radius=0.2)
label_box(s, x, 3.6, 3.25, 0.8, ["Genetic algorithm", ("NSGA-II", 14, False, GRAY)], fill=WHITE, line=ORANGE, size=18, radius=0.2)
label_box(s, x, 4.6, 3.25, 0.8, "One shared repair operator", fill=TINT, size=16, bold=True, radius=0.2)
x = xs[2] + 0.3
text(s, x, 2.55, 3.3, 0.9, "0.70 vs 0.25", 40, True, GREEN, "c")
text(s, x, 3.4, 3.3, 0.6, "front quality, optimisers vs random search", 14, False, DARK, "c")
text(s, x, 4.15, 3.3, 0.9, "1.34 to 0.53", 40, True, ORANGE, "c")
text(s, x, 5.0, 3.3, 0.6, "Sharpe ratio, training to held-out data", 14, False, DARK, "c")
text(s, 0.65, 6.15, 12.2, 0.55, "A good front is not a promise of future return.", 24, True, DARK, "c", italic=True)
notes(s, "Portfolio choice trades return against volatility and tail loss. We built two optimisers on one shared constraint-repair step. Both beat random search clearly, but portfolios picked for the best Sharpe ratio lost most of their advantage on unseen data.")

# ------------------------------------------------------------------ 3 Intro: problem
s = S[3]; prep(s, "Introduction: the problem")
im = Image.open(FIG + "fig_fronts_main.png"); w_, h_ = im.size
im.crop((0, 0, int(w_ * 0.505), h_)).save("crop_front_a.png")
pic(s, "crop_front_a.png", 0.75, 1.5, w=5.4, name="Efficient frontier figure")
text(s, 0.6, 6.2, 6.0, 0.4, "Each dot is a feasible portfolio; the front is the best trade-off", 13, False, GRAY, "c")
goals = [("Maximise annual return", BLUE), ("Minimise volatility", DARK), ("Minimise CVaR: average loss on the worst 5% of days", GRAY)]
for i, (g, c) in enumerate(goals):
    y = 1.6 + i * 1.15
    box(s, 7.1, y, 5.7, 0.95, TINT, name=f"Goal {i+1}")
    circle(s, 7.3, y + 0.2, 0.55, c, "")
    text(s, 8.05, y + 0.08, 4.65, 0.8, g, 18, True if i < 2 else False, DARK, anchor="m")
text(s, 7.1, 5.1, 5.7, 0.4, "Constraints applied to every portfolio", 16, True, BLUE)
chips = [("Long-only", 1.45), ("Fully invested", 1.75), ("Cap 30% per stock", 2.05), ("Min holding 1%", 1.85), ("At most K holdings", 2.1)]
pos_ = [(7.1, 5.6), (8.67, 5.6), (10.55, 5.6), (7.1, 6.15), (9.07, 6.15)]
for (c_, wd), (px, py) in zip(chips, pos_):
    label_box(s, px, py, wd, 0.45, c_, fill=WHITE, line=LINE, size=14, radius=0.4)
notes(s, "A portfolio has to balance return against volatility and tail loss, and real mandates add constraints. The result is not one answer but a front of trade-offs.")

# ------------------------------------------------------------------ 4 Intro: what we built
s = S[4]; prep(s, "Introduction: what we built")
tiles = [("Python library", "Portfolio model, repair operator, automated tests", BLUE),
         ("Two optimisers", "MOPSO and NSGA-II on one feasible space", ORANGE),
         ("Fair evaluation", "Hypervolume, random search, 30 held-out markets", GREEN),
         ("Web workbench", "Watch the search and inspect any portfolio", DARK)]
for i, (h, d, c) in enumerate(tiles):
    x = 0.65 + i * 3.1
    box(s, x, 1.65, 2.9, 3.3, TINT2, name=f"Built tile {i+1}")
    circle(s, x + 1.0, 1.9, 0.9, c, h[0], 30)
    text(s, x + 0.15, 3.0, 2.6, 0.5, h, 20, True, DARK, "c")
    text(s, x + 0.2, 3.55, 2.5, 1.3, d, 16, False, DARK, "c")
text(s, 0.65, 5.3, 4, 0.4, "Sustainable Development Goals", 16, True, BLUE)
for i, g in enumerate(["SDG 4  Quality Education", "SDG 8  Decent Work and Economic Growth", "SDG 9  Industry, Innovation, Infrastructure"]):
    label_box(s, 0.65 + [0, 3.2, 8.2][i], 5.8, [3.0, 4.9, 4.4][i], 0.5, g, fill=WHITE, line=LINE, size=14, radius=0.4)
text(s, 0.65, 6.4, 12, 0.4, "Scope: simulated markets, long-only portfolios, no trading costs in the optimisation", 14, False, GRAY)
notes(s, "Four deliverables: the library, two optimisers, a fair evaluation protocol and a browser workbench. Scope is limited to simulated markets, which is stated as a limitation.")

# ------------------------------------------------------------------ 5 Literature
s = S[5]; prep(s, "Literature Survey")
line_y = 3.45
box(s, 0.8, line_y - 0.03, 11.8, 0.06, GRAY, shape=MSO_SHAPE.RECTANGLE, name="Timeline")
tl = [("1952", "Markowitz", "mean-variance"), ("1995", "Kennedy, Eberhart", "particle swarm"), ("2000", "Rockafellar; Chang", "CVaR; cardinality search"),
      ("2002", "Deb et al.", "NSGA-II"), ("2004", "Coello et al.", "multi-objective PSO"), ("2009-12", "Cura; Zhu; Deng", "PSO for portfolios"), ("2012", "Metaxiotis", "MOEA review")]
for i, (yr, who, what) in enumerate(tl):
    cx = 1.35 + i * 1.75
    circle(s, cx - 0.14, line_y - 0.14, 0.28, BLUE if i % 2 == 0 else ORANGE, "")
    above = i % 2 == 0
    yy = 1.75 if above else 3.8
    text(s, cx - 0.95, yy, 1.9, 1.5, [(yr, 20, True, BLUE if above else ORANGE), (who, 14, True, DARK), (what, 14, False, DARK)], align="c", anchor="b" if above else "t")
text(s, 0.65, 5.2, 6, 0.4, "Research gaps we address", 16, True, BLUE)
for i, g in enumerate(["Held-out tests over many markets", "One repair operator for both algorithms", "Three-objective swarm front with CVaR", "Search you can watch and inspect"]):
    label_box(s, 0.65 + i * 3.1, 5.7, 2.95, 0.85, g, fill=TINT, size=14, radius=0.15)
notes(s, "From Markowitz to NSGA-II and multi-objective PSO. The surveyed portfolio-PSO papers optimise one scalarised objective and evaluate mostly in-sample, which motivates the four gaps shown.")

# ------------------------------------------------------------------ 6 Requirements
s = S[6]; prep(s, "Requirement Analysis and Specification")
stats = [("12", "functional requirements", BLUE), ("7", "non-functional requirements", ORANGE), ("0", "INR software cost, all open source", GREEN)]
for i, (n, l, c) in enumerate(stats):
    x = 0.65 + i * 4.15
    box(s, x, 1.6, 3.85, 2.3, TINT2, name=f"Stat {i+1}")
    text(s, x, 1.7, 3.85, 1.2, n, 66, True, c, "c")
    text(s, x + 0.2, 3.0, 3.45, 0.8, l, 18, False, DARK, "c")
text(s, 0.65, 4.3, 6, 0.4, "Software stack", 16, True, BLUE)
for i, (c_, wd) in enumerate([("Python 3.13", 1.9), ("NumPy", 1.4), ("pandas", 1.4), ("Matplotlib", 1.7), ("SciPy", 1.3), ("pytest", 1.3), ("Browser canvas", 2.2)]):
    xs_ = 0.65 + sum([1.9, 1.4, 1.4, 1.7, 1.3, 1.3, 2.2][:i]) + 0.12 * i
    label_box(s, xs_, 4.8, wd, 0.55, c_, fill=WHITE, line=LINE, size=15, radius=0.4)
box(s, 0.65, 5.7, 12.15, 0.85, TINT, name="Performance")
text(s, 0.85, 5.75, 11.8, 0.75, "Default run: 100 individuals, 150 iterations, 20 assets finishes in about 1.5 seconds", 20, True, DARK, "c", "m")
notes(s, "Requirements cover data input, constraints, both optimisers, live progress, inspection and error messages. Everything runs on free software; a default run takes about a second and a half.")

# ------------------------------------------------------------------ 7 Design: architecture
s = S[7]; prep(s, "Design: architecture")
bands = [(1.5, 1.35, TINT, "Browser", ["Control panel", "Live swarm canvas", "Portfolio inspector", "Wealth and convergence charts"]),
         (3.35, 1.35, "E8F6F1", "Server", ["POST /api/run", "GET /api/jobs/id", "POST .../cancel", "Job manager and worker threads"]),
         (5.2, 1.45, TINT2, "Optimisation core", ["Portfolio model and repair", "MOPSO", "NSGA-II", "Pareto utilities", "Baselines and back-test"])]
for y, h, fill, name, boxes in bands:
    box(s, 0.65, y, 12.15, h, fill, name=f"Band {name}")
    text(s, 0.8, y + 0.05, 3, 0.35, name, 14, True, BLUE if name == "Browser" else (GREEN if name == "Server" else DARK))
    n = len(boxes); gap = 0.2; bw = (11.65 - gap * (n - 1)) / n
    for i, b in enumerate(boxes):
        label_box(s, 0.9 + i * (bw + gap), y + 0.48, bw, h - 0.65, b, fill=WHITE, line=LINE, size=15, radius=0.15)
arrow(s, 6.7, 2.85, 6.7, 3.35); arrow(s, 6.7, 3.35, 6.7, 2.85)
arrow(s, 6.7, 4.7, 6.7, 5.2); arrow(s, 6.7, 5.2, 6.7, 4.7)
text(s, 6.85, 2.9, 2.5, 0.4, "JSON over HTTP", 13, False, GRAY)
text(s, 6.85, 4.75, 2.5, 0.4, "function calls", 13, False, GRAY)
notes(s, "Three layers. The browser draws and sends JSON. The server runs jobs on worker threads and stores the latest snapshot of each population. The core is plain Python and is shared by the tests, the command line and the web app.")

# ------------------------------------------------------------------ 8 Design: methodology
s = S[8]; prep(s, "Design: methodology")
text(s, 0.65, 1.45, 12, 0.4, "One repair operator, shared by both algorithms", 18, True, BLUE)
steps = ["Clip negatives to zero", "Keep the K largest", "Drop tiny weights", "Add assets if the cap needs more", "Cap at 30% and redistribute", "Normalise to sum 1"]
for i, st in enumerate(steps):
    x = 0.65 + i * 2.0
    sp = box(s, x, 1.95, 2.15, 1.1, BLUE if i in (0, 5) else TINT, shape=MSO_SHAPE.CHEVRON, name=f"Repair step {i+1}")
    sp.adjustments[0] = 0.25
    tf = sp.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.35); tf.margin_right = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].text = st; para_fmt(tf.paragraphs[0], 13, True, WHITE if i in (0, 5) else DARK, False, PP_ALIGN.CENTER)
text(s, 0.65, 3.3, 12, 0.4, "Six experiments", 18, True, BLUE)
ex = [("E1", "Front quality", "10 seeds against random search"), ("E2", "Holding limits", "K = none, 10, 8 and 5"), ("E3", "Number of assets", "10 to 60 assets"),
      ("E4", "Search budget", "5 to 200 iterations"), ("E5", "Held-out test", "30 simulated markets"), ("E6", "Operator ablation", "with and without")]
for i, (c_, h, d) in enumerate(ex):
    r, cc = divmod(i, 3)
    x, y = 0.65 + cc * 4.15, 3.8 + r * 1.35
    box(s, x, y, 3.85, 1.2, TINT2, name=f"Experiment {c_}")
    circle(s, x + 0.15, y + 0.3, 0.6, BLUE, c_, 16)
    text(s, x + 0.9, y + 0.12, 2.85, 0.5, h, 18, True, DARK)
    text(s, x + 0.9, y + 0.62, 2.85, 0.5, d, 14, False, DARK)
text(s, 0.65, 6.5, 12, 0.4, "Seeded factor-model markets with fat tails and a crisis period; 70% training, 30% held-out", 14, False, GRAY)
notes(s, "The repair operator guarantees every portfolio is long-only, fully invested, capped and within the holding limit. Six experiments vary one factor at a time, with the last using 30 independent markets.")

# ------------------------------------------------------------------ 9 Implementation: optimisers
s = S[9]; prep(s, "Implementation: the two optimisers")
cfg = [(0.65, BLUE, "MOPSO (particle swarm)", ["Particles start as feasible portfolios", "Leader chosen from the archive by crowding", "Velocity update, then repair", "Archive keeps the non-dominated set"], ["100 particles", "inertia 0.9 to 0.4", "c1 = c2 = 1.5"]),
       (6.85, ORANGE, "NSGA-II (genetic algorithm)", ["Population of feasible portfolios", "Tournament by rank and crowding", "SBX crossover, mutation and swap", "Parents and children: keep the best"], ["100 individuals", "SBX index 15", "swap mutation 0.2"])]
for x, c, h, st, ch_ in cfg:
    label_box(s, x, 1.5, 5.8, 0.65, h, fill=c, color=WHITE, size=20, bold=True, radius=0.3)
    for i, t_ in enumerate(st):
        label_box(s, x, 2.35 + i * 0.85, 5.8, 0.68, t_, fill=TINT if c == BLUE else "FBEFE3", size=16, radius=0.15)
        if i < 3: arrow(s, x + 2.9, 3.03 + i * 0.85, x + 2.9, 3.2 + i * 0.85, c)
    for j, p_ in enumerate(ch_):
        label_box(s, x + j * 1.97, 5.85, 1.85, 0.5, p_, fill=WHITE, line=c, size=13, radius=0.4)
text(s, 0.65, 6.45, 12.2, 0.4, "Same budget for both: 15,100 objective evaluations per run", 16, True, DARK, "c")
notes(s, "Both algorithms see exactly the same problem and the same evaluation budget. The swarm uses an external archive and crowding-based leaders; the genetic algorithm is standard NSGA-II with a swap mutation added.")

# ------------------------------------------------------------------ 10 Implementation: workbench
s = S[10]; prep(s, "Implementation: the workbench")
pic(s, FIG + "ui_overview.png", 0.6, 1.55, w=7.6, name="Workbench screenshot")
cards = [("Watch the search", "Both populations converge on their fronts, live", BLUE), ("Inspect any portfolio", "Click a point: weights, training and held-out results", ORANGE), ("Bring your own prices", "Upload a CSV, set constraints, press Run", GREEN)]
for i, (h, d, c) in enumerate(cards):
    y = 1.55 + i * 1.7
    box(s, 8.55, y, 4.25, 1.5, TINT2, name=f"Callout {i+1}")
    circle(s, 8.75, y + 0.45, 0.6, c, "")
    text(s, 9.5, y + 0.12, 3.2, 0.5, h, 18, True, DARK)
    text(s, 9.5, y + 0.62, 3.2, 0.85, d, 14, False, DARK)
notes(s, "The browser workbench runs both optimisers side by side and animates them. After the run you can click any point on a front to see its weights and how it behaved on held-out data.")

# ------------------------------------------------------------------ 11 Testing
s = S[11]; prep(s, "Testing and Validation")
for i, (n, l, c) in enumerate([("14", "automated tests pass", BLUE), ("19", "system test cases: server and browser", ORANGE), ("0", "constraint violations in every experiment run", GREEN)]):
    x = 0.65 + i * 4.15
    box(s, x, 1.6, 3.85, 2.2, TINT2, name=f"Test stat {i+1}")
    text(s, x, 1.65, 3.85, 1.2, n, 66, True, c, "c")
    text(s, x + 0.2, 2.85, 3.45, 0.85, l, 16, False, DARK, "c")
lv = [("Component", "Repair, objectives, Pareto maths"), ("Algorithm", "Feasible, non-dominated, reproducible, beats random"), ("System", "Server jobs, error paths, browser, phone, dark mode")]
for i, (h, d) in enumerate(lv):
    x = 0.65 + i * 4.3
    label_box(s, x, 4.3, 3.7, 0.65, h, fill=[BLUE, ORANGE, GREEN][i], color=WHITE, size=20, bold=True, radius=0.3)
    label_box(s, x, 5.1, 3.7, 1.3, d, fill=WHITE, line=LINE, size=16, radius=0.1)
    if i < 2: arrow(s, x + 3.75, 4.62, x + 4.25, 4.62)
notes(s, "Tests run at three levels. Fourteen automated tests, nineteen system test cases including browser checks, and zero constraint violations across all experiment runs. Same seed gives the same front.")

# ------------------------------------------------------------------ 12 Results: front quality
s = S[12]; prep(s, "Results: front quality")
pic(s, FIG + "fig_fronts_main.png", 0.55, 1.5, w=7.7, name="Fronts figure")
hv = F["e1"]["hv"]
cd = CategoryChartData(); cd.categories = ["MOPSO", "NSGA-II", "Random"]; cd.add_series("Hypervolume", [round(hv["PSO"][0], 3), round(hv["GA"][0], 3), round(hv["Random"][0], 3)])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(8.5), Inches(1.45), Inches(4.3), Inches(3.4), cd); gf.name = "Hypervolume chart"
ch = gf.chart; style_chart(ch, 13, legend=False); chart_title(ch, "Hypervolume (higher is better)")
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.number_format = "0.00"; pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.data_labels.font.size = Pt(14); pl.data_labels.font.bold = True
ch.value_axis.maximum_scale = 0.9; ch.value_axis.minimum_scale = 0; ch.value_axis.tick_labels.number_format = "0.0"; ch.value_axis.tick_labels.number_format_is_linked = False
for i, c in enumerate([BLUE, ORANGE, GRAY]):
    p = pl.series[0].points[i]; p.format.fill.solid(); p.format.fill.fore_color.rgb = rgb(c)
for i, (n, l) in enumerate([("100 vs 20", "non-dominated portfolios per run: optimisers vs random"), ("p = 0.002", "both optimisers vs random, paired test over 10 seeds"), ("0.70 vs 0.71", "swarm vs genetic: a statistical tie in practice")]):
    x = 0.65 + i * 4.15
    box(s, x, 5.1, 3.85, 1.5, TINT2, name=f"Result stat {i+1}")
    text(s, x, 5.15, 3.85, 0.7, n, 32, True, [BLUE, GREEN, ORANGE][i], "c")
    text(s, x + 0.2, 5.85, 3.45, 0.7, l, 14, False, DARK, "c")
notes(s, "Both optimisers find roughly one hundred non-dominated portfolios against about twenty for random search, with hypervolume about 0.70 versus 0.25. The genetic algorithm is marginally ahead on average but the difference is tiny next to the seed-to-seed spread.")

# ------------------------------------------------------------------ 13 Results: budget
s = S[13]; prep(s, "Results: search budget")
e4 = F["e4"]
cd = CategoryChartData(); cd.categories = [r[1] for r in e4["table"]]
cd.add_series("MOPSO", [round(v, 3) for v in e4["hv"]["PSO"]]); cd.add_series("NSGA-II", [round(v, 3) for v in e4["hv"]["GA"]]); cd.add_series("Random search", [round(v, 3) for v in e4["hv"]["Random"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.5), Inches(8.0), Inches(5.0), cd); gf.name = "Budget chart"
ch = gf.chart; style_chart(ch, 14); chart_title(ch, "Hypervolume against objective evaluations", 16)
for ser, c in zip(ch.plots[0].series, [BLUE, ORANGE, GRAY]):
    ser.format.line.color.rgb = rgb(c); ser.format.line.width = Pt(3); ser.smooth = False
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(c); ser.marker.format.line.color.rgb = rgb(c); ser.marker.size = 8
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1.0; ch.value_axis.tick_labels.number_format = "0.0"; ch.value_axis.tick_labels.number_format_is_linked = False
cards = [("Up to 50 iterations", "Swarm ahead in 5 of 5 seeds", BLUE), ("From 100 iterations", "Genetic algorithm ahead in 5 of 5 seeds", ORANGE), ("Random search", "Stays far behind at every budget", GRAY)]
for i, (h, d, c) in enumerate(cards):
    y = 1.55 + i * 1.7
    box(s, 8.95, y, 3.85, 1.5, TINT2, name=f"Budget callout {i+1}")
    circle(s, 9.15, y + 0.45, 0.6, c, "")
    text(s, 9.9, y + 0.15, 2.85, 0.5, h, 18, True, DARK)
    text(s, 9.9, y + 0.65, 2.85, 0.8, d, 14, False, DARK)
notes(s, "With few evaluations the swarm is clearly better; it reaches in 10 iterations what the genetic algorithm needs 25 to 50 iterations to match. From 100 iterations the genetic algorithm finishes slightly ahead.")

# ------------------------------------------------------------------ 14 Results: held-out
s = S[14]; prep(s, "Results: held-out performance")
sm = F["e5"]["sharpe_means"]
keys = [("Equal weight", "Equal weight"), ("PSO best Sharpe", "Swarm: best Sharpe"), ("GA best Sharpe", "GA: best Sharpe"), ("PSO min volatility", "Swarm: min vol."), ("GA min volatility", "GA: min vol.")]
cd = CategoryChartData(); cd.categories = [k[1] for k in keys]
cd.add_series("Training window", [round(sm[k[0]][0], 2) for k in keys]); cd.add_series("Held-out window", [round(sm[k[0]][1], 2) for k in keys])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.45), Inches(6.3), Inches(4.0), cd); gf.name = "Sharpe chart"
ch = gf.chart; style_chart(ch, 12); chart_title(ch, "Sharpe ratio: training vs held-out")
pl = ch.plots[0]; pl.gap_width = 70; pl.overlap = -5; pl.has_data_labels = True; pl.data_labels.number_format = "0.00"; pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.data_labels.font.size = Pt(12)
color_series(pl.series[0], MID); color_series(pl.series[1], DARK)
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1.6; ch.value_axis.tick_labels.number_format = "0.0"; ch.value_axis.tick_labels.number_format_is_linked = False
def pctv(row): return float(row[4].replace("%", ""))
tb = {r[0]: r for r in F["e5"]["table"]}
vol = [pctv(tb["Equal weight"]), pctv(tb["MOPSO best Sharpe"]), pctv(tb["NSGA-II best Sharpe"]), pctv(tb["MOPSO min volatility"]), pctv(tb["NSGA-II min volatility"])]
cd = CategoryChartData(); cd.categories = [k[1] for k in keys]; cd.add_series("Held-out volatility (%)", vol)
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.1), Inches(1.45), Inches(5.7), Inches(4.0), cd); gf.name = "Volatility chart"
ch = gf.chart; style_chart(ch, 12, legend=False); chart_title(ch, "Held-out volatility (% a year)")
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.number_format = "0.0"; pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.data_labels.font.size = Pt(12)
for i, c in enumerate([GREEN, BLUE, ORANGE, BLUE, ORANGE]):
    p = pl.series[0].points[i]; p.format.fill.solid(); p.format.fill.fore_color.rgb = rgb(c)
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 20
box(s, 0.65, 5.7, 12.15, 0.95, DARK, name="Takeaway banner")
text(s, 0.85, 5.72, 11.75, 0.9, "Return does not transfer to new data. Risk does: lower held-out volatility in 30 of 30 markets", 22, True, WHITE, "c", "m")
notes(s, "Average over 30 independent simulated markets. Portfolios chosen for the best Sharpe ratio looked excellent in training but fell below equal weight on unseen data. Minimum-volatility portfolios were calmer than equal weight in every one of the 30 markets.")

# ------------------------------------------------------------------ 15 Results: robustness
s = S[15]; prep(s, "Results: robustness checks")
for i in range(3):
    box(s, 0.65 + i * 4.15, 1.55, 3.85, 4.9, TINT2, name=f"Robustness tile {i+1}")
text(s, 0.65, 1.65, 3.85, 0.5, "Scale", 20, True, BLUE, "c")
sec = F["e3"]["sec"]
cd = CategoryChartData(); cd.categories = ["10", "20", "40", "60"]; cd.add_series("MOPSO", [round(v, 2) for v in sec["PSO"]]); cd.add_series("NSGA-II", [round(v, 2) for v in sec["GA"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(2.2), Inches(3.55), Inches(3.0), cd); gf.name = "Runtime chart"
ch = gf.chart; style_chart(ch, 12); ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 3
ch.value_axis.tick_labels.number_format = "0"; ch.value_axis.tick_labels.number_format_is_linked = False
color_series(ch.plots[0].series[0], BLUE); color_series(ch.plots[0].series[1], ORANGE); ch.plots[0].gap_width = 60
text(s, 0.8, 5.3, 3.55, 1.0, ["Seconds per run, 10 to 60 assets", ("about 1.5 s whatever the size", 14, True, DARK)], 13, False, DARK, "c")
text(s, 4.8, 1.65, 3.85, 0.5, "Constraints", 20, True, GREEN, "c")
text(s, 4.8, 2.5, 3.85, 1.5, "0", 96, True, GREEN, "c")
text(s, 5.0, 4.4, 3.45, 1.8, "violations of the cap or the holding limit, for limits of none, 10, 8 and 5 holdings", 16, False, DARK, "c")
text(s, 8.95, 1.65, 3.85, 0.5, "Operator ablation", 20, True, ORANGE, "c")
text(s, 8.95, 2.7, 3.85, 1.2, "No gain", 44, True, ORANGE, "c")
text(s, 9.15, 4.0, 3.45, 2.2, ["No turbulence: p = 0.64", "No swap mutation: p = 0.55", ("Neither changes front quality at these settings", 14, True, DARK)], 16, False, DARK, "c", spacing=10)
notes(s, "Run time stays near one and a half seconds from 10 to 60 assets. The repair operator kept every portfolio feasible. The two extra operators we added did not measurably help, which we report honestly.")

# ------------------------------------------------------------------ 16 Conclusion
s = S[16]; prep(s, "Conclusion")
cc = [("Both optimisers work", "Far better fronts than random search; every constraint respected", BLUE, "1"), ("Return does not transfer", "Best-Sharpe picks: 1.34 in training, 0.53 held out; equal weight 0.59", ORANGE, "2"), ("Risk does", "Minimum-volatility picks were calmer than equal weight in 30 of 30 markets", GREEN, "3")]
for i, (h, d, c, n) in enumerate(cc):
    x = 0.65 + i * 4.15
    box(s, x, 1.6, 3.85, 3.4, TINT2, name=f"Takeaway {i+1}")
    circle(s, x + 1.4, 1.85, 1.05, c, n, 36)
    text(s, x + 0.15, 3.1, 3.55, 0.6, h, 22, True, DARK, "c")
    text(s, x + 0.25, 3.75, 3.35, 1.2, d, 16, False, DARK, "c")
text(s, 0.65, 5.3, 6, 0.4, "Future scope", 16, True, BLUE)
for i, (g, w) in enumerate([("Real market data", 2.7), ("Shrinkage estimators", 3.0), ("Trading costs as an objective", 3.5), ("Hosted classroom version", 2.6)]):
    xx = 0.65 + sum([2.7, 3.0, 3.5, 2.6][:i]) + 0.12 * i
    label_box(s, xx, 5.8, w, 0.6, g, fill=WHITE, line=LINE, size=15, radius=0.4)
notes(s, "Three take-aways: the optimisers work, historical returns do not predict future returns well, but risk estimates do transfer. Future work: real market data, shrinkage estimators, costs as an objective and a hosted version for classes.")

# ------------------------------------------------------------------ 17 POs
s = S[17]; prep(s, "Program Outcomes Attained")
pos = [("PO1", "Engineering knowledge", 3), ("PO2", "Problem analysis", 3), ("PO3", "Design of solutions", 2), ("PO4", "Investigation of problems", 3), ("PO5", "Modern tool usage", 3), ("PO6", "Engineer and the world", 1), ("PO7", "Ethics", 2),
       ("PO8", "Teamwork", 3), ("PO9", "Communication", 2), ("PO10", "Project management", 2), ("PO11", "Life-long learning", 2), ("PSO1", "Design AI systems", 3), ("PSO2", "Decisions with AI and data", 2), ("PSO3", "End-to-end systems", 3)]
fills = {3: (BLUE, WHITE), 2: (MID, DARK), 1: ("DCE6EE", DARK)}
for i, (c_, n, lv) in enumerate(pos):
    r, cc_ = divmod(i, 7)
    x, y = 0.65 + cc_ * 1.75, 1.95 + r * 1.85
    f, tc = fills[lv]
    sp = label_box(s, x, y, 1.62, 1.7, [(c_, 20, True, tc), (n, 13, False, tc), (f"Level {lv}", 13, True, tc)], fill=f, size=13, radius=0.1)
    sp.name = f"{c_} tile"
for i, (lab, lv) in enumerate([("3  High", 3), ("2  Medium", 2), ("1  Low", 1)]):
    label_box(s, 9.4 + i * 1.2, 1.45, 1.1, 0.38, lab, fill=fills[lv][0], color=fills[lv][1], size=12, bold=True, radius=0.4)
text(s, 0.65, 5.75, 12, 0.5, "Each level is justified in the project report (attainment sheet).", 14, False, GRAY)
notes(s, "Levels as assessed in the report. Strongest: engineering knowledge, problem analysis, investigations, tool usage and teamwork. Weakest: the engineer and the world, since societal impact is only addressed through the SDGs.")

prs.save("SwarmFolio_Final_Presentation.pptx")
print("saved", len(prs.slides), "slides")
